import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from ite_core import demo
from ite_core import session as S
from modules import data_export as X

ALL = [
    "Read me",
    "Residents",
    "ITE by exam",
    "By PGY level",
    "By class",
    "Year over year",
    "YoY summary",
    "MKSAP",
    "MKSAP sections",
    "Crosswalk",
    "Item gaps",
    "Area ranks",
]


@pytest.fixture(scope="module")
def ds():
    d = demo.make_demo()
    r = S.load_rosters([S.Upload(n, b) for n, b in d["rosters"]])
    m = S.load_mksap(S.Upload(*d["mksap"]))
    cw = S.match(r, m)
    return S.Dataset(rosters=r, mksap=m, crosswalk=cw, wide=S.build_wide(r, m, cw))


def _read(b):
    return pd.read_excel(io.BytesIO(b), sheet_name=None)


def test_full_workbook(ds):
    x = _read(X.build_workbook(ds, demo.sample_program_report()))
    assert list(x) == ALL
    assert len(x["Residents"]) == 30 and len(x["ITE by exam"]) == 90
    assert {"Graduating class", "Points from Low risk", "MKSAP questions answered"} <= set(x["Residents"].columns)
    assert "Cardiology percentile" in x["ITE by exam"].columns
    assert x["Read me"].Description.str.contains("resident-identifiable").any()
    assert x["Area ranks"].columns[0] == "Content area" and "PGY-1 program rank" in x["Area ranks"].columns


def test_graduating_class_is_consistent(ds):
    x = _read(X.build_workbook(ds, None))
    e = x["ITE by exam"]
    assert (e["Graduating class"] == e["Exam year"] + 4 - e["PGY"]).all()


def test_sheets_follow_uploads(ds):
    no_mk = S.Dataset(rosters=ds.rosters, mksap=None, crosswalk=None, wide=S.build_wide(ds.rosters, None, None))
    assert "MKSAP" not in X.included_sheets(no_mk, None) and "Item gaps" not in X.included_sheets(no_mk, None)
    assert X.included_sheets(None, demo.sample_program_report()) == ["Item gaps", "Area ranks"]
    assert X.included_sheets(None, None) == []


def test_overview_export_has_no_section_sheet(ds):
    from tests.test_mksap_tab import _overview_from_detail

    d = demo.make_demo()
    mk = S.load_mksap(S.Upload("overview.csv", _overview_from_detail(d["mksap"][1])))
    cw = S.match(ds.rosters, mk)
    ov = S.Dataset(rosters=ds.rosters, mksap=mk, crosswalk=cw, wide=S.build_wide(ds.rosters, mk, cw))
    sheets = X.included_sheets(ov, None)
    assert "MKSAP" in sheets and "MKSAP sections" not in sheets


def test_workbook_is_styled_and_filterable(ds):
    wb = load_workbook(io.BytesIO(X.build_workbook(ds, None)))
    ws = wb["Residents"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws["A1"].font.bold


def test_cutoff_changes_tiers(ds):
    strict = _read(X.build_workbook(ds, None, cutoffs=(0.5, 0.95)))["Residents"]
    loose = _read(X.build_workbook(ds, None, cutoffs=(0.5, 0.75)))["Residents"]
    assert (strict["Risk tier 2026"] == "Low").sum() <= (loose["Risk tier 2026"] == "Low").sum()
