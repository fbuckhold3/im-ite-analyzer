"""Combined program data export: one identified Excel workbook with every table the app computes.

Pure function: build_workbook(dataset, program, cutoffs) -> bytes. A sheet is included only when its data
was uploaded. Row-level and identified — never offered in a de-identified form (rows re-identify easily).
"""

from __future__ import annotations

import datetime as dt
import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ite_core import board, history, mksap, pdf_reports, stats
from ite_core.nomogram import DEFAULT_CUTOFFS
from ite_core.roster import AREA_LABELS

from .items_tab import item_frame

HEADER_FILL = PatternFill("solid", fgColor="E7F0FB")


def _ite_by_exam(long: pd.DataFrame) -> pd.DataFrame:
    rename = {
        "ID": "ACP Candidate ID",
        "Name": "Resident",
        "Exam_Year": "Exam year",
        "PGY": "PGY",
        "Grad_Class": "Graduating class",
        "Pct_Correct": "Total % correct",
        "Percentile": "Total percentile",
        "Pass_Prob": "Estimated ABIM pass probability",
        "Risk_Tier": "Risk tier (default cutoffs)",
        "Program_Number": "ACP program number",
        "Program_Name": "Program",
    }
    for key, label in AREA_LABELS.items():
        rename[f"pc_{key}"] = f"{label} % correct"
        rename[f"ile_{key}"] = f"{label} percentile"
    cols = [c for c in rename if c in long]
    return long[cols].rename(columns=rename).sort_values(["Graduating class", "Resident", "Exam year"])


def _residents(w: pd.DataFrame, cutoffs) -> pd.DataFrame:
    bt = board.board_table(w, cutoffs)
    cy, years = bt.attrs["current_year"], bt.attrs["years"]
    out = pd.DataFrame(
        {"ACP Candidate ID": bt.ID, "Resident": bt.Name, f"PGY ({cy})": bt.PGY, "Graduating class": cy + 4 - bt.PGY}
    )
    for y in years:
        out[f"% correct {y}"] = bt[f"Pct_{y}"]
    out["Change in % correct (last year)"] = bt.dPct
    for y in years:
        out[f"Percentile {y}"] = bt[f"Ile_{y}"]
    out["Change in percentile (last year)"] = bt.dIle
    for y in years:
        out[f"Pass estimate {y}"] = bt[f"Prob_{y}"].round(3)
    for y in years:
        out[f"Risk tier {y}"] = bt[f"Tier_{y}"]
    out["Tier change"] = bt.Tier_change
    out["% correct for Low risk at current PGY"] = bt.Needed_Low
    out["Points from Low risk"] = bt.Gap_Low
    if "mk_answered" in bt:
        out["MKSAP questions answered"] = bt.mk_answered
        out["MKSAP % correct"] = bt.mk_pct_correct.round(1)
    return out.sort_values([f"PGY ({cy})", "Resident"], ascending=[False, True])


def _agg(df: pd.DataFrame, keys: dict) -> pd.DataFrame:
    rename = {
        **keys,
        "n": "n",
        "pc": "Mean % correct",
        "ile": "Mean percentile",
        "med": "Median percentile",
        "ge50": "% at/above 50th percentile",
        "lt35": "% below 35th percentile",
        "pp_low": "% Low risk",
        "small_n": "Fewer than 5 residents",
    }
    return df[[c for c in rename if c in df]].rename(columns=rename)


def _sheets(ds, program, cutoffs) -> list[tuple[str, pd.DataFrame, str]]:
    """(sheet name, table, description) for every sheet with data."""
    out = []
    if ds is not None and ds.wide is not None:
        long, w = ds.rosters.long, ds.wide
        out.append(
            (
                "Residents",
                _residents(w, cutoffs),
                f"One row per resident in the {ds.rosters.current_year} roster: scores by year, estimates, tiers, "
                "changes, MKSAP totals.",
            )
        )
        out.append(
            (
                "ITE by exam",
                _ite_by_exam(long),
                "Every resident x exam year (long format), with all content-area scores. Best for your own analysis.",
            )
        )
        out.append(
            (
                "By PGY level",
                _agg(history.by_level(long, cutoffs), {"PGY": "PGY", "Exam_Year": "Exam year"}),
                "Each PGY level in each exam year (different residents in each row).",
            )
        )
        out.append(
            (
                "By class",
                _agg(
                    history.by_class(long, cutoffs),
                    {"Grad_Class": "Graduating class", "PGY": "PGY", "Exam_Year": "Exam year"},
                ),
                "Each graduating class at each PGY level.",
            )
        )
        pts = stats.paired_points(w)
        if not pts.empty:
            out.append(
                (
                    "Year over year",
                    pts.rename(
                        columns={
                            "ID": "ACP Candidate ID",
                            "Name": "Resident",
                            "PGY": "Current PGY",
                            "pc_prev": "% correct last year",
                            "pc_cur": "% correct this year",
                            "dPct_1yr": "Change in % correct",
                            "ile_prev": "Percentile last year",
                            "ile_cur": "Percentile this year",
                            "dPctile_1yr": "Change in percentile",
                        }
                    ),
                    "Same residents, last year -> this year.",
                )
            )
            summ = []
            for p in (2, 3):
                c = stats.paired_change(w, p)
                if c["n"]:
                    summ.append(
                        {
                            "Current PGY": p,
                            "n": c["n"],
                            "Mean percentile change": c["d_ile"],
                            "Median percentile change": c["d_med"],
                            "Mean % correct change": c["d_pc"],
                            "Rose 10+ percentile points": c["up"],
                            "Fell 10+ percentile points": c["down"],
                            "Wilcoxon p (percentile)": c["p"],
                        }
                    )
            out.append(("YoY summary", pd.DataFrame(summ), "Paired change by PGY with Wilcoxon signed-rank test."))
        if ds.mksap is not None:
            ppl = ds.mksap.people
            link = (
                ds.crosswalk[["mksap_key", "acp_id", "ite_name"]].dropna(subset=["mksap_key"])
                if ds.crosswalk is not None
                else pd.DataFrame(columns=["mksap_key", "acp_id", "ite_name"])
            )
            mk = ppl.merge(link, left_on="mk_key", right_on="mksap_key", how="left")
            out.append(
                (
                    "MKSAP",
                    pd.DataFrame(
                        {
                            "MKSAP last name": mk.mk_last,
                            "MKSAP first name": mk.mk_first,
                            "Linked ACP Candidate ID": mk.acp_id,
                            "Linked resident": mk.ite_name,
                            "Questions answered": mk.mk_answered,
                            "Correct": mk.mk_correct,
                            "% correct": mk.mk_pct_correct.round(1),
                            "Questions in bank": mk.mk_bank,
                            "% of bank": mk.mk_pct_bank.round(1),
                        }
                    ).sort_values("MKSAP last name"),
                    f"Every account in the {mksap.KIND_LABEL[ds.mksap.kind]} export; unlinked accounts are usually "
                    "graduates or faculty.",
                )
            )
            if ds.mksap.sections is not None and "mk_answered" in w:
                st = mksap.section_table(ds.mksap.sections, w)
                st = st[["name", "Subspecialty", "PGY", "bank", "a", "c", "n", "done", "acc"]]
                out.append(
                    (
                        "MKSAP sections",
                        st.rename(
                            columns={
                                "Subspecialty": "Section code",
                                "name": "Section",
                                "PGY": "PGY",
                                "a": "Answered",
                                "c": "Correct",
                                "n": "Question slots (size x residents)",
                                "bank": "Questions in section",
                                "done": "% attempted",
                                "acc": "Accuracy %",
                            }
                        ),
                        "Section coverage by current PGY (full detail export only).",
                    )
                )
        if ds.crosswalk is not None:
            out.append(
                (
                    "Crosswalk",
                    ds.crosswalk.rename(
                        columns={
                            "acp_id": "ACP Candidate ID",
                            "ite_name": "Resident",
                            "mksap_key": "MKSAP key",
                            "mksap_last": "MKSAP last name",
                            "mksap_first": "MKSAP first name",
                            "tier": "Match",
                            "n_candidates": "Candidates found",
                            "confirmed_by_user": "Reviewed by you",
                        }
                    ),
                    "How each resident was matched to an MKSAP account.",
                )
            )
    if program is not None:
        out.append(
            (
                "Item gaps",
                item_frame(program["items"]),
                "Every item on the Program Performance Report with program vs national % correct and flags.",
            )
        )
        rank_names = {"area": "Content area"}
        for k, label in (("1", "PGY-1"), ("2", "PGY-2"), ("3", "PGY-3"), ("A", "All PGYs")):
            rank_names |= {
                f"all{k}": f"{label} all programs %",
                f"our{k}": f"{label} program %",
                f"rank{k}": f"{label} program rank",
            }
        ranks = pdf_reports.area_ranks(program["areas"]).rename(columns=rank_names)
        out.append(("Area ranks", ranks, "Program percentile rank among all programs, by content area and PGY group."))
    return out


def _style(ws, df: pd.DataFrame) -> None:
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200).tolist()])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))


def included_sheets(ds, program) -> list[str]:
    return [name for name, _, _ in _sheets(ds, program, DEFAULT_CUTOFFS)]


def build_workbook(ds, program, cutoffs=DEFAULT_CUTOFFS) -> bytes:
    sheets = _sheets(ds, program, cutoffs)
    files = [f for f in (ds.rosters.files.file.tolist() if ds is not None and len(ds.rosters.files) else [])]
    if ds is not None and ds.mksap is not None:
        files.append(ds.mksap.name)
    if program is not None:
        files.append(program.get("file", "Program Performance Report"))
    readme = pd.DataFrame(
        [
            ("CONFIDENTIAL", "Contains resident-identifiable education records. Do not forward or post."),
            ("Exported", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Low-risk cutoff", f"Estimated ABIM pass probability >= {int(cutoffs[1] * 100)}% (High risk < 50%)"),
            ("Source files", "; ".join(files)),
            (
                "Pass estimates",
                "Logistic approximation of McDonald et al., Acad Med 2020 (doi:10.1097/ACM.0000000000003382)",
            ),
            ("", ""),
        ]
        + [(name, desc) for name, _, desc in sheets]
        + [("", ""), ("Created with", "im-ite-analyzer — github.com/fbuckhold3/im-ite-analyzer (Apache-2.0)")],
        columns=["Item", "Description"],
    )
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="Read me", index=False)
        _style(xw.sheets["Read me"], readme)
        xw.sheets["Read me"].column_dimensions["B"].width = 100
        for name, df, _ in sheets:
            df.to_excel(xw, sheet_name=name[:31], index=False)
            _style(xw.sheets[name[:31]], df)
    return buf.getvalue()
